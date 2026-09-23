"""B-006: list -> node -> counterparty -> back -> node XLSX, with real data."""
from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest
from streamlit.testing.v1 import AppTest

from ui.card import counterparties
from ui.data import load_bundle, load_raw
from ui.theme import LABELS

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("GRAF_OUT", str(ROOT / "out"))
    monkeypatch.setenv("GRAF_DATA", str(ROOT / "data"))
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("OPENAI_MODEL", "")
    result = AppTest.from_file(ROOT / "app.py", default_timeout=30).run()
    assert not result.exception
    return result


def _button(app, label):
    return next(item for item in app.button if item.label == label)


def _top_table(app):
    return next(item for item in app.dataframe if item.key and item.key.startswith("top_"))


def _select_row(app, dataframe, index):
    # Streamlit 1.64's AppTest Dataframe has no row-selection setter. Send the
    # same string_value widget event as the browser to exercise the real callback.
    states = app._tree.get_widget_states()
    state = next((item for item in states.widgets if item.id == dataframe.proto.id), None)
    if state is None:
        state = states.widgets.add()
        state.id = dataframe.proto.id
    state.string_value = json.dumps({"selection": {"rows": [index], "columns": [], "cells": []}})
    app._run(states)
    assert not app.exception


def _open_gid(app, gid):
    app.text_input(key="search").set_value(gid).run()
    _button(app, "Открыть узел").click().run()
    assert not app.exception
    assert app.session_state["gid"] == gid
    assert app.radio(key="page").value == "Узел"


def _downloads(app):
    return [item for item in app.get("download_button") if item.proto.label == "Скачать XLSX узла"]


def _assert_node_workbook(app, expected_gid):
    gid, payload = app.session_state["xlsx_node"]
    assert gid == expected_gid and isinstance(gid, str) and len(gid) == 18
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["Сводка"]
    headers = [cell.value for cell in sheet[1]]
    gid_column = headers.index("gid") + 1
    cells = [sheet.cell(row, gid_column) for row in range(2, sheet.max_row + 1)]
    assert [cell.value for cell in cells] == [expected_gid]
    assert all(cell.data_type == "s" and cell.number_format == "@" for cell in cells)
    limits = list(workbook["Границы данных"].values)
    positions = {name: index for index, name in enumerate(limits[0])}
    requests = [row for row in limits[1:] if row[positions["тип"]] == "запрос"]
    assert all(row[positions["gid"]] == expected_gid for row in requests)
    expected = load_bundle(ROOT / "out").tables["data_requests"]
    assert len(requests) == expected.gid.eq(int(expected_gid)).sum()
    downloads = _downloads(app)
    assert len(downloads) == 1
    assert downloads[0].proto.id.endswith(f"card_download_{expected_gid}")


@pytest.mark.parametrize("direction,peer_column", [("in", "src"), ("out", "dst")])
def test_counterparty_selection_and_back_keep_exact_gid_and_filters(app, direction, peer_column):
    roles = ["coordinator", "consolidator"]
    app.radio(key="analysis_mode").set_value("Деньги курьеров").run()
    app.slider[0].set_value(7).run()
    app.multiselect(key="role_filter").set_value(roles).run()
    original_gid = str(_top_table(app).value["gid"].iloc[0])
    _select_row(app, _top_table(app), 0)
    assert app.session_state["gid"] == original_gid
    assert app.radio(key="page").value == "Узел"

    _, _, transactions = load_raw(ROOT / "data")
    peer_table = counterparties(transactions, int(original_gid), direction)
    expected_gid = str(peer_table[peer_column].iloc[0])
    frame = next(item for item in app.dataframe if item.key == f"peer_{original_gid}_{direction}")
    _select_row(app, frame, 0)
    assert app.session_state["gid"] == expected_gid
    assert isinstance(app.session_state["gid"], str)
    assert app.query_params["gid"] == [expected_gid]
    assert app.radio(key="analysis_mode").value == "Деньги курьеров"
    assert app.slider[0].value == 7
    assert app.multiselect(key="role_filter").value == roles

    _button(app, "← К предыдущему узлу").click().run()
    assert not app.exception
    assert app.session_state["gid"] == original_gid
    assert app.query_params["gid"] == [original_gid]
    assert app.multiselect(key="role_filter").value == roles
    assert app.radio(key="analysis_mode").value == "Деньги курьеров"
    assert app.slider[0].value == 7
    # Revisiting the same peer after Back must work, not retain a stale row event.
    frame = next(item for item in app.dataframe if item.key == f"peer_{original_gid}_{direction}")
    _select_row(app, frame, 0)
    assert app.session_state["gid"] == expected_gid
    _button(app, "← К предыдущему узлу").click().run()
    _button(app, "← К топ-листу").click().run()
    assert not app.exception
    assert app.radio(key="page").value == "Топ-лист"
    assert app.multiselect(key="role_filter").value == roles
    assert app.radio(key="analysis_mode").value == "Деньги курьеров"
    assert app.slider[0].value == 7


def test_card_xlsx_is_bound_to_current_gid_and_never_reuses_previous_node(app):
    gids = load_bundle(ROOT / "out").tables["top_nodes"].gid.astype(str).head(2).tolist()
    first, second = gids
    _open_gid(app, first)
    app.button(key=f"card_xlsx_{first}").click().run()
    assert not app.exception and app.radio(key="page").value == "Узел"
    _assert_node_workbook(app, first)

    _open_gid(app, second)
    assert not _downloads(app)
    assert app.button(key=f"card_xlsx_{second}")
    app.button(key=f"card_xlsx_{second}").click().run()
    assert not app.exception
    _assert_node_workbook(app, second)

    _button(app, "← К предыдущему узлу").click().run()
    assert not app.exception and app.session_state["gid"] == first
    assert not _downloads(app)
    assert app.button(key=f"card_xlsx_{first}")


def test_hidden_node_keeps_filters_until_explicit_show_action(app):
    gid = str(load_bundle(ROOT / "out").tables["top_nodes"].iloc[0].gid)
    _open_gid(app, gid)
    app.multiselect(key="role_filter").set_value(["payer"]).run()
    _button(app, "Показать на карте").click().run()
    assert not app.exception
    assert app.multiselect(key="role_filter").value == ["payer"]
    assert any("Выбранный узел скрыт" in item.value for item in app.info)
    assert app.session_state["gid"] == gid
    app.button(key="reveal_node").click().run()
    assert not app.exception
    assert set(app.multiselect(key="role_filter").value) == set(LABELS)
    assert app.radio(key="analysis_mode").value == "Структура"
    assert app.session_state["gid"] == gid
    assert not any("Выбранный узел скрыт" in item.value for item in app.info)
    assert any(f"Выбран узел {gid}" in item.value for item in app.caption)


def test_empty_top_explains_filters_and_reset_restores_saved_data(app):
    app.multiselect(key="role_filter").set_value(["payer"]).run()
    assert not app.exception and _top_table(app).value.empty
    assert any("В топ-30 нет узлов с такими условиями" in item.value for item in app.info)
    assert any("Данные в выгрузке есть" in item.value for item in app.info)
    assert any("Активные фильтры" in item.value for item in app.caption)
    app.button(key="reset_visible").click().run()
    assert not app.exception
    assert set(app.multiselect(key="role_filter").value) == set(LABELS)
    assert app.radio(key="analysis_mode").value == "Структура"
    assert len(_top_table(app).value) == len(load_bundle(ROOT / "out").tables["top_nodes"])


def test_changing_sources_invalidates_old_xlsx_and_navigation_history(app, tmp_path):
    first, second = load_bundle(ROOT / "out").tables["top_nodes"].gid.astype(str).head(2)
    _open_gid(app, first)
    app.button(key=f"card_xlsx_{first}").click().run()
    _open_gid(app, second)
    assert app.session_state["node_history"]
    source = next(item for item in app.text_input if item.label == "Исходные данные: папка или ZIP")
    source.set_value(str(tmp_path / "other-data")).run()
    assert not app.exception
    assert "xlsx_node" not in app.session_state
    assert not app.session_state.get("node_history", [])
    assert not any(item.label == "← К предыдущему узлу" for item in app.button)
    assert not _downloads(app)

