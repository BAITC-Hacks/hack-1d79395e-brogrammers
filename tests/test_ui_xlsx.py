"""Evidence export: precision, provenance, per-node scope and graceful A2 input."""
from io import BytesIO
from pathlib import Path
import shutil
import pandas as pd
import pytest
from tools.ui_fixtures import ensure_stub
from openpyxl import load_workbook
from graf.evidence_xlsx import SHEETS, evidence_frames, workbook_bytes, build_evidence, build_node_evidence, transaction_evidence, expand_paths
from ui.data import load_bundle


@pytest.mark.parametrize('source', ['out_stub', 'out'])
def test_workbook_contract_and_precision(source):
    frames = evidence_frames('data', source)
    workbook = load_workbook(BytesIO(workbook_bytes(frames)))
    assert workbook.sheetnames == SHEETS
    for sheet in workbook:
        assert sheet.freeze_panes == 'A2'
        headers = [c.value for c in sheet[1]]
        for column, name in enumerate(headers, 1):
            if name not in {'gid','src','dst','seed_gid','target_gid'}:
                continue
            for row in sheet.iter_rows(min_row=2):
                cell = row[column-1]
                if cell.value is None or cell.value == 'пороги':
                    continue
                assert cell.data_type == 's' and cell.number_format == '@'
                assert isinstance(cell.value, str) and len(cell.value) == 18 and cell.value.isdigit()
    summary_ids = [r[1].value for r in workbook['Сводка'].iter_rows(min_row=2)]
    assert summary_ids == load_bundle(source).tables['top_nodes'].sort_values('rank').head(30).gid.astype(str).tolist()


def test_one_node_and_build_api(tmp_path):
    bundle = load_bundle('out_stub')
    gid = str(bundle.tables['top_nodes'].iloc[0].gid)
    wb = load_workbook(BytesIO(build_node_evidence(gid, 'data', 'out_stub')))
    assert wb['Сводка'].max_row == 2
    sheet = wb['Транзакции-основания']
    assert all(gid in (row[0].value, row[1].value) for row in sheet.iter_rows(min_row=2))
    output = tmp_path/'out_stub'
    shutil.copytree('out_stub', output)
    assert build_evidence('data', output) is None
    assert load_workbook(output/'evidence.xlsx').sheetnames == SHEETS
    assert not (output/'evidence.tmp.xlsx').exists()
    with pytest.raises(ValueError, match='отсутствует'):
        build_node_evidence('999999999999999999', 'data', 'out_stub')


def test_transaction_reasons_and_literal_text():
    ids = [10**17+i for i in range(6)]
    tx = pd.DataFrame({'src':[ids[0],ids[1],ids[2],ids[3]], 'dst':[ids[3],ids[3],ids[3],ids[4]],
                       'date':pd.to_datetime(['2026-07-01']*3+['2026-07-03']), 'sum_kzt':[5000]*4})
    nodes = pd.DataFrame({'gid':ids, 'role':['payer']*3+['consolidator','terminal','peripheral']})
    paths = pd.DataFrame(columns=['target_gid','seed_gid','hops','path_rank','path_gids','path_days','path_amounts','bottleneck_kzt'])
    result = transaction_evidence(tx,nodes,{ids[3]},paths)
    assert all('синхронный сбор' in s and 'сбор' in s for s in result.iloc[:3]['подтверждает'])
    assert result.iloc[-1]['подтверждает'] == 'пересылка ≤2 дней'
    wb = load_workbook(BytesIO(workbook_bytes({'literal':pd.DataFrame({'evidence':['=1+1']})})))
    assert wb['literal']['A2'].data_type == 's'
    assert wb['literal']['A2'].value == '=1+1'


def test_full_workbook_keeps_all_requests_but_node_export_stays_scoped():
    bundle = load_bundle('out')
    expected = bundle.tables['data_requests']
    frames = evidence_frames('data', 'out')
    requests = frames['Границы данных'].loc[lambda frame: frame['тип'].eq('запрос')]
    actual = {(row.gid, row.описание, row.причина) for row in requests.itertuples()}
    expected_rows = {(str(row.gid), row.request, row.reason) for row in expected.itertuples()}
    assert len(requests) == len(expected)
    assert actual == expected_rows
    top_ids = set(bundle.tables['top_nodes'].gid.astype(str))
    outside_top = next(gid for gid, _, _ in expected_rows if gid not in top_ids)
    node_limits = evidence_frames('data', 'out', outside_top)['Границы данных']
    node_requests = node_limits.loc[node_limits['тип'].eq('запрос')]
    assert set(node_requests.gid) == {outside_top}
    assert len(node_requests) == int(expected.gid.eq(int(outside_top)).sum())


def test_role_criteria_use_runtime_gates_and_explain_seed_exception(monkeypatch):
    from graf import config
    from graf.evidence_xlsx import role_criteria
    monkeypatch.setattr(config, 'CONS_MIN_IN_DEG', 9)
    monkeypatch.setattr(config, 'COORD_MIN_KEY_LINKS', 4)
    monkeypatch.setattr(config, 'PAYER_MAX_KZT', 75000)
    criteria = role_criteria().set_index('роль')
    assert 'in_deg ≥ 9' in criteria.loc['consolidator', 'ворота']
    assert 'from_key ≥ 4 и to_key ≥ 4' in criteria.loc['coordinator', 'ворота']
    assert 'out_kzt ≤ 75000' in criteria.loc['payer', 'ворота']
    terminal = criteria.loc['terminal', 'ворота']
    assert 'depth ≤ 3; in_deg ≥ 1' in terminal
    assert 'не seed и pass_through < 0.1' in terminal
    assert '1−p_forward ≥ 0.6' in terminal
    assert 'не seed' in criteria.loc['transit', 'ворота']
    assert 'max(0,1−|1−pass_through|/0.2)' in criteria.loc['transit', 'скор']
    assert 'coordinator → payer' in criteria.loc['Порядок выбора', 'ворота']
