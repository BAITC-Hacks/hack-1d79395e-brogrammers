"""Analyst navigation and fallback behavior against the real published outputs."""
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pytest
from streamlit.testing.v1 import AppTest

from ui.data import load_bundle
from ui.navigation import selected_gid
from ui.theme import LABELS

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / 'app.py'


@pytest.fixture
def real_app(monkeypatch):
    monkeypatch.setenv('GRAF_OUT', str(ROOT / 'out'))
    monkeypatch.setenv('GRAF_DATA', str(ROOT / 'data'))
    # Empty environment values also keep load_dotenv from loading live keys.
    monkeypatch.setenv('OPENAI_API_KEY', '')
    monkeypatch.setenv('OPENAI_MODEL', '')
    app = AppTest.from_file(APP, default_timeout=30).run()
    assert not app.exception
    return app


def _button(app, label):
    return next(button for button in app.button if button.label == label)


def _open_node(app, gid):
    app.text_input(key='search').set_value(gid).run()
    _button(app, 'Открыть узел').click().run()
    assert not app.exception
    assert app.session_state['gid'] == gid
    assert app.radio(key='page').value == 'Узел'


def test_map_selection_preserves_exact_18_digit_id():
    # Adjacent IDs above 2**53 must stay distinct through selection.
    first, second = str(10**17 + 1), str(10**17 + 2)
    event = {'selection': {'points': [{'customdata': [second, 'роль', 0.1]}]}}
    assert selected_gid(event, [first, second]) == second
    assert event['selection']['points'][0]['customdata'][0] == second
    event['selection']['points'].insert(0, {'customdata': ['unknown']})
    assert selected_gid(event, [first, second]) == second


@pytest.mark.parametrize('custom', [None, [], {}, [10**17 + 1], [float(10**17 + 1)], ['999999999999999999']])
def test_map_selection_rejects_missing_rounded_or_unknown_ids(custom):
    event = {'selection': {'points': [{'customdata': custom}]}}
    assert selected_gid(event, [str(10**17 + 1)]) is None
    assert selected_gid({}, [str(10**17 + 1)]) is None


def test_real_node_map_evidence_workflow_preserves_selected_gid(real_app):
    app = real_app
    bundle = load_bundle(ROOT / 'out')
    gid = str(bundle.tables['top_nodes'].sort_values('rank').iloc[1].gid)
    _open_node(app, gid)
    assert any(info.value == bundle.nodes.loc[bundle.nodes.gid.astype(str).eq(gid), 'evidence'].iloc[0]
               for info in app.info)
    _button(app, 'Показать на карте').click().run()
    assert not app.exception
    assert app.radio(key='page').value == 'Карта по коленам'
    assert app.session_state['map_selected'] == gid
    assert app.session_state['gid'] == gid
    assert any(f'Выбран узел {gid}' in caption.value for caption in app.caption)
    _button(app, f'Открыть узел {gid}').click().run()
    _button(app, 'Доказательства узла').click().run()
    assert not app.exception
    assert app.radio(key='page').value == 'Доказательства'
    assert app.session_state['gid'] == gid
    assert any(gid in info.value for info in app.info)
    _button(app, f'Подготовить XLSX узла {gid}').click().run()
    assert not app.exception
    selected, payload = app.session_state['xlsx_node']
    assert selected == gid
    sheet = load_workbook(BytesIO(payload))['Сводка']
    headers = [cell.value for cell in sheet[1]]
    column = headers.index('gid') + 1
    assert [sheet.cell(row, column).value for row in range(2, sheet.max_row + 1)] == [gid]


def test_empty_top_filters_can_be_reset_in_one_action(real_app):
    app = real_app
    app.radio(key='analysis_mode').set_value('Деньги курьеров').run()
    app.multiselect(key='role_filter').set_value([]).run()
    assert not app.exception
    assert any('В топ-30 нет узлов' in info.value for info in app.info)
    _button(app, 'Показать весь топ-лист').click().run()
    assert not app.exception
    assert app.radio(key='analysis_mode').value == 'Структура'
    assert set(app.multiselect(key='role_filter').value) == set(LABELS)
    table = next(frame.value for frame in app.dataframe if {'gid', 'rank', 'why'} <= set(frame.value.columns))
    assert len(table) == len(load_bundle(ROOT / 'out').tables['top_nodes'])


def test_money_layer_uses_saved_counts_without_raw_data(real_app, tmp_path):
    app = real_app
    source = next(text for text in app.text_input if text.label == 'Исходные данные: папка или ZIP')
    source.set_value(str(tmp_path / 'missing-data')).run()
    app.radio(key='analysis_mode').set_value('Деньги курьеров').run()
    nodes = load_bundle(ROOT / 'out').nodes
    for gap, column in [(2, 'seed_exp_fast'), (31, 'seed_exp_chrono')]:
        app.select_slider[0].set_value(gap).run()
        assert not app.exception
        expected = int((nodes[column].gt(0) & ~nodes.is_seed).sum())
        assert any(f'Узлов в следе {expected} из' in info.value for info in app.info)
        assert not any('Денежный слой недоступен' in warning.value for warning in app.warning)
        expected_filtered = int((nodes[column].gt(0) | nodes.is_seed).sum())
        metric = next(metric for metric in app.metric if metric.label == 'Узлов после фильтров')
        assert metric.value == str(expected_filtered)
