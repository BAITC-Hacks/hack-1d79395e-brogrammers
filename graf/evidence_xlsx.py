"""B3: six-sheet evidence export. Does not assign roles or infer new money flows."""
from io import BytesIO
import importlib
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from ui.data import load_bundle, load_raw, display_frame, ID_COLUMNS
from ui.theme import LABELS, LIMITS, DISCLAIMER
from ui.criteria import role_criteria, ROLE_ORDER

SHEETS = ["Сводка", "Признаки", "Транзакции-основания", "Пути денег", "Критерии", "Границы данных"]


def expand_paths(paths):
    """path_days may specify transfer days (hops) or node days (hops+1)."""
    rows = []
    for r in paths.to_dict("records"):
        ids = str(r["path_gids"]).split(">")
        days = str(r["path_days"]).split(">")
        amounts = str(r["path_amounts"]).split(">")
        if len(days) == len(ids):
            days = days[1:]
        if len(ids) != int(r["hops"])+1 or len(days) != len(ids)-1 or len(amounts) != len(days):
            raise ValueError(f"paths.csv: несогласованный путь к {r['target_gid']}")
        for step, (src, dst, day, amount) in enumerate(zip(ids, ids[1:], days, amounts), 1):
            if not 1 <= int(day) <= 31:
                raise ValueError("paths.csv: день должен быть в диапазоне 1–31")
            rows.append(dict(target_gid=str(r["target_gid"]), seed_gid=str(r["seed_gid"]), path_rank=r["path_rank"],
                             step=step, src=src, dst=dst, day=int(day), sum_kzt=float(amount), bottleneck_kzt=r["bottleneck_kzt"]))
    return pd.DataFrame(rows, columns=["target_gid", "seed_gid", "path_rank", "step", "src", "dst", "day", "sum_kzt", "bottleneck_kzt"])


def config_rows():
    try:
        config = importlib.import_module("graf.config")
    except ModuleNotFoundError as exc:
        if exc.name != "graf.config":
            raise
        return [{"параметр": "graf.config", "значение": "Модуль A ещё не передан; пороги не подтверждены"}]
    rows = []
    for name in sorted(vars(config)):
        value = getattr(config, name)
        if name.isupper() and isinstance(value, (str, int, float, bool, list, tuple, dict, set)):
            rows.append({"параметр": name, "значение": json.dumps(value, ensure_ascii=False, default=str)})
    return rows or [{"параметр": "graf.config", "значение": "Нет экспортированных констант"}]


def transaction_evidence(tx, nodes, selected_ids, paths):
    subset = tx.loc[tx.src.isin(selected_ids) | tx.dst.isin(selected_ids)].copy()
    roles = nodes.set_index("gid").role.to_dict()
    sync = tx.groupby(["dst", "date"]).src.nunique().to_dict()
    incoming_days = {int(gid): set(group.date) for gid, group in tx.groupby("dst")}
    steps = expand_paths(paths)
    path_keys = {(int(r.src), int(r.dst), int(r.day)) for r in steps.itertuples()}
    tags = []
    for r in subset.itertuples():
        reasons = []
        if roles.get(int(r.dst)) in {"consolidator", "coordinator"}:
            reasons.append("сбор")
        if sync.get((r.dst, r.date), 0) >= 3:
            reasons.append("синхронный сбор")
        if any(r.date-pd.Timedelta(days=d) in incoming_days.get(int(r.src), set()) for d in range(3)):
            reasons.append("пересылка ≤2 дней")
        if (int(r.src), int(r.dst), int(r.date.day)) in path_keys:
            reasons.append("маршрут от seed")
        if roles.get(int(r.src)) == "distributor":
            reasons.append("веерная рассылка")
        tags.append("; ".join(reasons) or "контекст")
    subset["подтверждает"] = tags
    return subset[["src", "dst", "date", "sum_kzt", "подтверждает"]].sort_values(["date", "src", "dst"])


def evidence_frames(data_dir, out_dir, gid=None):
    bundle = load_bundle(out_dir)
    _, _, tx = load_raw(data_dir)
    nodes = bundle.nodes
    top = bundle.tables["top_nodes"].sort_values("rank").head(30)
    if gid is None:
        ids = set(top.gid.astype(int))
        summary = top.copy()
    else:
        ids = {int(gid)}
        summary = nodes.loc[nodes.gid.eq(int(gid)), [c for c in ["rank", "gid", "role", "priority_score", "evidence"] if c in nodes]].copy()
        if summary.empty:
            raise ValueError(f"Узел {gid} отсутствует в выгрузке")
        summary = summary.rename(columns={"evidence": "why"})
    extra = [c for c in ["role_label", "confidence_level", "visibility", "counter_signals"] if c in nodes and c not in summary]
    summary = summary.merge(nodes[["gid"]+extra], on="gid", how="left", validate="one_to_one")
    features = nodes.loc[nodes.gid.isin(ids)].copy()
    config = config_rows()
    # A labelled final row records the exact available runtime thresholds.
    threshold_row = {"gid": "пороги", "evidence": json.dumps(config, ensure_ascii=False)}
    features = pd.concat([display_frame(features), pd.DataFrame([threshold_row])], ignore_index=True)
    paths = bundle.tables["paths"]
    paths = paths.loc[paths.target_gid.isin(ids)]
    steps = expand_paths(paths)
    transactions = transaction_evidence(tx, nodes, ids, paths)
    criteria = pd.DataFrame(role_criteria()+[dict(роль="Порядок решения", ворота=ROLE_ORDER)])
    criteria = pd.concat([criteria, pd.DataFrame(config).rename(columns={"параметр": "роль", "значение": "ворота"}).assign(источник="graf.config: текущие параметры")], ignore_index=True)
    limits = [{"тип": "ограничение", "описание": text} for text in [DISCLAIMER]+LIMITS]
    limits.append({"тип": "источник", "описание": f"Выгрузки: {Path(out_dir).name}; транзакции: {Path(data_dir).name}"})
    if bundle.is_stub:
        limits.insert(0, {"тип": "ЗАГЛУШКА", "описание": "Роли, скоры, денежный след и устойчивость — демонстрационные; не использовать как доказательства."})
    if bundle.missing:
        limits.append({"тип": "неполная выгрузка", "описание": ", ".join(bundle.missing)})
    requests = bundle.tables["data_requests"]
    if gid is not None:
        requests = requests.loc[requests.gid.isin(ids)]
    for r in requests.itertuples():
        limits.append({"тип": "запрос", "gid": str(r.gid), "описание": r.request, "причина": r.reason})
    return dict(zip(SHEETS, [display_frame(summary), features, display_frame(transactions), steps, criteria, pd.DataFrame(limits)]))


def workbook_bytes(frames):
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl", datetime_format="YYYY-MM-DD") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF", size=11)
                cell.fill = PatternFill("solid", fgColor="193F64")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[1].height = 32
            for column in sheet.columns:
                header = str(column[0].value)
                longest = max([len(str(cell.value or "")) for cell in column]+[len(header)])
                width = 23 if header in ID_COLUMNS else min(70, max(14, longest+2))
                sheet.column_dimensions[column[0].column_letter].width = width
                for cell in column[1:]:
                    if isinstance(cell.value, str):
                        # Keep literal evidence strings, including an initial '=', as text.
                        cell.data_type = "s"
                    if header in ID_COLUMNS:
                        cell.number_format = "@"
                    elif isinstance(cell.value, float):
                        cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for row in sheet.iter_rows(min_row=2):
                lines = max((len(str(c.value or "")) // max(10,int(sheet.column_dimensions[c.column_letter].width)) + 1 for c in row), default=1)
                sheet.row_dimensions[row[0].row].height = min(300, max(30, 15*lines))
    return stream.getvalue()


def build_evidence(data_dir, out_dir) -> None:
    payload = workbook_bytes(evidence_frames(data_dir, out_dir))
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    temporary = Path(out_dir) / "evidence.tmp.xlsx"
    temporary.write_bytes(payload)
    temporary.replace(Path(out_dir) / "evidence.xlsx")


def build_node_evidence(gid, data_dir, out_dir) -> bytes:
    return workbook_bytes(evidence_frames(data_dir, out_dir, gid))
